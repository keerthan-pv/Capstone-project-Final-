# module5/document_upload.py
"""
Document Upload — Let users upload their own PDF/TXT/DOCX/CSV/MD and query it.
Creates a temporary in-memory vector store from the parsed content.

Security guards:
  • file size capped by configs.config.MAX_UPLOAD_BYTES
  • magic-byte sniffing + extension allow-list (defence in depth)
  • path-traversal-safe filename storage
"""

from __future__ import annotations

import logging
import os
import re
import sys
from pathlib import PurePosixPath
from typing import List, Optional

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from langchain_community.vectorstores import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from configs.config import CHUNK_OVERLAP, CHUNK_SIZE, MAX_UPLOAD_BYTES

logger = logging.getLogger(__name__)


# Magic-byte signatures for the file types we accept.
_MAGIC_BYTES = {
    "pdf":  (b"%PDF",),
    "docx": (b"PK\x03\x04",),  # DOCX is a ZIP container
}


def _safe_basename(name: str) -> str:
    """
    Reduce an uploaded filename to a safe basename.
    Strips any directory components and drops control characters.
    """
    base = PurePosixPath(name.replace("\\", "/")).name
    return "".join(c for c in base if 32 <= ord(c) < 127) or "uploaded"


class UploadValidationError(ValueError):
    """Raised when an uploaded file fails security validation."""


def normalize_upload_query(query: str) -> str:
    """Normalize common chat shorthand before vector retrieval.

    Uploaded documents often use spaced headings (``Programming Languages``),
    whereas users naturally type camel-cased or abbreviated forms.
    """
    text = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", query.strip())
    text = re.sub(r"\b(?:r|u)\b", lambda match: {"r": "are", "u": "you"}[match.group(0).lower()], text, flags=re.IGNORECASE)
    # Expand common resume/document concepts before retrieval.  This is query
    # expansion—not a fabricated answer—and only helps locate source text.
    aliases = {
        "college": "college university institution education",
        "university": "university college institution education",
        "youtube": "youtube channel tutorials video",
        "channel": "channel youtube tutorials",
        "tutorial": "tutorial youtube channel",
        "programming": "programming coding technical skills",
        "coding": "coding programming technical skills",
        "framework": "framework technology technical skills",
        "frameworks": "frameworks framework technology technical skills",
        "technology": "technology framework technical skills",
        "technologies": "technologies technology framework technical skills",
        "job": "job experience employment work",
        "company": "company employer experience work",
        "employer": "employer company experience work",
    }
    alias_pattern = r"\b(?:" + "|".join(map(re.escape, aliases)) + r")\b"
    text = re.sub(
        alias_pattern,
        lambda match: aliases[match.group(0).casefold()],
        text,
        flags=re.IGNORECASE,
    )
    return re.sub(r"\s+", " ", text).strip()


def _terms(text: str) -> set[str]:
    """Tokenise text using the same shorthand/camel-case normalisation."""
    return set(re.findall(r"[a-z0-9]+", normalize_upload_query(text).casefold()))


def extract_document_answer(query: str, documents: List[Document]) -> str:
    """Return an answer using only retrieved uploaded-document text.

    Upload questions must not fall back to web search: a document can contain
    personal details that are unavailable elsewhere.  A small structured pass
    handles common document requests such as a resume's project list, then the
    existing local extractive QA implementation handles general questions.
    """
    context = "\n\n".join(
        document.page_content.strip()
        for document in documents
        if document.page_content and document.page_content.strip()
    )
    if not context:
        return "I could not find relevant text in the uploaded document."

    normalized = normalize_upload_query(query).casefold()

    # Compact answers for labelled resume fields.  These patterns use the
    # uploaded evidence directly and prevent a whole heading line from being
    # returned when the user requests one value.
    if re.search(r"\b(?:youtube|channel)\b", normalized):
        channel = re.search(
            r"\|\s*([^|\n]+?)\s*-\s*\d+(?:k|m)?\+?\s*youtube\b",
            context,
            re.IGNORECASE,
        )
        if channel:
            return f"The YouTube channel is {channel.group(1).strip()}."

    if re.search(r"\b(?:email|e-mail)\b", normalized):
        email = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", context)
        if email:
            return f"The email address is {email.group(0)}."

    if re.search(r"\b(?:college|university|institution)\b", normalized):
        institution = re.search(
            r"(?:EDUCATION\s+)?([A-Z][A-Za-z& .]+(?:University|College|Institute)[^\n]*)",
            context,
            re.IGNORECASE,
        )
        if institution:
            return institution.group(1).strip().rstrip(".")

    if re.search(r"\b(?:cgpa|gpa|grade|score)\b", normalized):
        grade = re.search(r"\b(?:CGPA|GPA)\b[^\n]*|\b\d+(?:\.\d+)?\s*(?:CGPA|GPA)\b", context, re.I)
        if grade:
            return grade.group(0).strip()

    if re.search(r"\b(?:degree|qualification)\b", normalized):
        degree = re.search(
            r"\b(?:Bachelor|Master|B\.?(?:E|Tech|Sc)|M\.?(?:E|Tech|Sc))\b[^\n]*",
            context,
            re.I,
        )
        if degree:
            return degree.group(0).strip().rstrip(".")

    if re.search(r"\b(?:programming|coding)\b.*\b(?:language|languages)\b|\b(?:language|languages)\b.*\b(?:programming|coding)\b", normalized):
        programming = re.search(r"Programming\s+Languages\s*:\s*([^\n]+)", context, re.I)
        if programming:
            return f"Programming languages: {programming.group(1).strip().rstrip('.')}"

    if re.search(r"\b(?:framework|technology|technologies|technical skills)\b", normalized):
        technologies = re.search(r"Technologies/Frameworks\s*:\s*([^\n]+(?:\n(?![A-Z][A-Z ]{3,}\n)[^\n]+)*)", context, re.I)
        if technologies:
            value = re.sub(r"\s+", " ", technologies.group(1)).strip().rstrip(".")
            return f"Technologies/frameworks: {value}"
    if re.search(r"\b(?:project|projects|portfolio)\b", normalized):
        # Resumes commonly use a compact statement such as "Other projects
        # include ...". Prefer it over an incidental mention of "project
        # management", which previously caused the wrong answer here.
        project_match = re.search(
            r"\b(?:other\s+)?projects?\s+(?:include|included|are)\s*[:\-]?\s*"
            r"(.+?)(?=\n[A-Z][A-Z\s]{3,}\n|$)",
            context,
            re.IGNORECASE | re.DOTALL,
        )
        if project_match:
            projects = re.sub(r"\s+", " ", project_match.group(1)).strip(" ,;.")
            if projects:
                return f"The document lists these projects: {projects}."

    # Do not call the generic RAG prompt fallback here. Its parser was written
    # for a prompt string and could mistake its own "SHORT, DIRECT answer"
    # instruction for source text.  Uploads instead use a transparent,
    # evidence-only sentence ranker over the chunks we just retrieved.
    stop_words = {
        "a", "an", "and", "are", "be", "did", "do", "does", "for", "from",
        "have", "he", "i", "in", "is", "it", "me", "of", "on", "please",
        "the", "their", "they", "this", "to", "was", "what", "which", "who",
        "with", "would", "you", "your",
    }
    query_terms = {
        token for token in _terms(query)
        if len(token) > 1 and token not in stop_words
    }
    if not query_terms:
        return "Please ask a more specific question about the uploaded document."

    unique_sentences: list[str] = []
    seen: set[str] = set()
    for sentence in re.split(r"(?<=[.!?])\s+|\n+", context):
        clean = re.sub(r"\s+", " ", sentence).strip(" -•\t")
        key = clean.casefold()
        if len(clean) >= 12 and key not in seen:
            seen.add(key)
            unique_sentences.append(clean)

    scored: list[tuple[int, str]] = []
    for sentence in unique_sentences:
        sentence_terms = _terms(sentence)
        overlap = len(query_terms & sentence_terms)
        if overlap:
            # Prefer concise evidence, but retain a long resume bullet when it
            # contains multiple query terms.
            score = overlap * 10 - min(len(sentence), 800) // 200
            scored.append((score, sentence))

    if not scored:
        return "I could not find an answer to that question in the uploaded document."

    scored.sort(key=lambda item: item[0], reverse=True)
    answer = scored[0][1]
    return answer[:700].rstrip()


class DocumentUploader:
    """
    Handles document upload, parsing, chunking, and indexing into
    a temporary vector store for querying.
    """

    SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".md", ".csv", ".docx"}

    def __init__(self, embedding_model=None, max_bytes: int = MAX_UPLOAD_BYTES):
        self.embedding_model = embedding_model
        self.max_bytes = max_bytes
        self._vectorstore: Optional[Chroma] = None
        self._chunks: List[Document] = []
        self._filename: str = ""

    @property
    def is_ready(self) -> bool:
        return self._vectorstore is not None and len(self._chunks) > 0

    @property
    def filename(self) -> str:
        return self._filename

    @property
    def num_chunks(self) -> int:
        return len(self._chunks)

    @property
    def vectorstore(self) -> Optional[Chroma]:
        """The in-memory vector store for the active document (None until indexed)."""
        return self._vectorstore

    @property
    def chunks(self) -> List[Document]:
        """Return a copy of the active document chunks for evidence-only fallback."""
        return list(self._chunks)

    # ── Validation ──────────────────────────────────────────────────────────
    def _validate(self, file_content: bytes, filename: str) -> str:
        if not file_content:
            raise UploadValidationError("Empty upload (0 bytes).")
        if len(file_content) > self.max_bytes:
            raise UploadValidationError(
                f"File too large: {len(file_content):,} bytes "
                f"(max allowed: {self.max_bytes:,} bytes)."
            )
        ext = os.path.splitext(filename)[1].lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise UploadValidationError(
                f"Unsupported file type '{ext}'. Supported formats: {sorted(self.SUPPORTED_EXTENSIONS)}"
            )

        # Magic-byte sniff for binary types within first 1KB (tolerates BOM/leading whitespace)
        header_sample = file_content[:1024]
        if ext == ".pdf" and b"%PDF" not in header_sample:
            raise UploadValidationError("File is not a valid PDF (missing %PDF signature header).")
        if ext == ".docx" and b"PK\x03\x04" not in header_sample:
            raise UploadValidationError("File is not a valid DOCX (missing ZIP/DOCX header).")
        return ext

    # ── Parsers ─────────────────────────────────────────────────────────────
    def parse_file(self, file_content: bytes, filename: str) -> List[Document]:
        ext = self._validate(file_content, filename)
        self._filename = _safe_basename(filename)

        if ext == ".pdf":
            return self._parse_pdf(file_content)
        if ext == ".docx":
            return self._parse_docx(file_content)
        return self._parse_text(file_content)

    def _parse_text(self, content: bytes) -> List[Document]:
        text = ""
        for encoding in ("utf-8", "latin-1", "cp1252", "utf-16"):
            try:
                text = content.decode(encoding)
                break
            except Exception:
                continue
        if not text:
            text = content.decode("utf-8", errors="replace")

        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        docs = [
            Document(
                page_content=para,
                metadata={"source": self._filename, "paragraph": i},
            )
            for i, para in enumerate(paragraphs)
            if len(para) > 10
        ]
        if not docs and text.strip():
            docs.append(Document(
                page_content=text.strip()[:10000],
                metadata={"source": self._filename, "paragraph": 0},
            ))
        if not docs:
            raise UploadValidationError("File contains no readable text.")
        return docs

    def _parse_pdf(self, content: bytes) -> List[Document]:
        docs: List[Document] = []
        parse_errors: List[str] = []

        # 1. Primary engine: PyMuPDF (fast, handles complex layouts, fonts, encodings)
        try:
            import pymupdf
            doc = pymupdf.open(stream=content, filetype="pdf")
            if doc.is_encrypted:
                try:
                    doc.authenticate("")
                except Exception:
                    pass

            for i, page in enumerate(doc):
                try:
                    text = page.get_text("text").strip()
                    if text:
                        docs.append(Document(
                            page_content=text,
                            metadata={"source": self._filename, "page": i + 1},
                        ))
                except Exception as page_exc:
                    logger.warning("PyMuPDF extraction failed on page %d: %s", i + 1, page_exc)
            doc.close()
        except Exception as exc:
            parse_errors.append(f"PyMuPDF error: {exc}")
            logger.warning("PyMuPDF parser failed, trying pypdf fallback: %s", exc)

        # 2. Secondary engine fallback: pypdf
        if not docs:
            try:
                import io
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(content))
                if reader.is_encrypted:
                    try:
                        reader.decrypt("")
                    except Exception:
                        pass

                for i, page in enumerate(reader.pages):
                    try:
                        text = (page.extract_text() or "").strip()
                        if text:
                            docs.append(Document(
                                page_content=text,
                                metadata={"source": self._filename, "page": i + 1},
                            ))
                    except Exception as page_exc:
                        logger.warning("pypdf page %d extraction failed: %s", i + 1, page_exc)
            except Exception as exc:
                parse_errors.append(f"pypdf error: {exc}")
                logger.warning("pypdf fallback failed: %s", exc)

        if not docs:
            err_msg = (
                "PDF contained no extractable text layer. "
                "This typically happens if the PDF consists of scanned images/photos without OCR. "
                "Please upload a text-searchable PDF or a document in .txt, .md, or .docx format."
            )
            if parse_errors:
                err_msg += f" (Diagnostics: {'; '.join(parse_errors)})"
            raise UploadValidationError(err_msg)

        return docs

    def _parse_docx(self, content: bytes) -> List[Document]:
        import io
        from docx import Document as DocxDocument
        try:
            doc = DocxDocument(io.BytesIO(content))
        except Exception as exc:
            raise UploadValidationError(f"Could not parse DOCX: {exc}") from exc

        parts = []
        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                parts.append(t)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        docs = [
            Document(
                page_content=p,
                metadata={"source": self._filename, "paragraph": i},
            )
            for i, p in enumerate(parts)
            if len(p) > 10
        ]
        if not docs:
            raise UploadValidationError("DOCX document contained no text content.")
        return docs

    # ── Indexing ────────────────────────────────────────────────────────────
    def index_documents(
        self,
        file_content: bytes,
        filename: str,
        chunk_size: int = CHUNK_SIZE,
        chunk_overlap: int = CHUNK_OVERLAP,
    ) -> int:
        docs = self.parse_file(file_content, filename)

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        self._chunks = splitter.split_documents(docs)

        if self.embedding_model is None:
            from module1.data_preprocessing import get_embedding_model
            self.embedding_model = get_embedding_model()

        # Replace any prior in-memory store. We do not pass persist_directory
        # so the collection lives only in memory and is GC'd when this object dies.
        self._vectorstore = Chroma.from_documents(
            documents=self._chunks,
            embedding=self.embedding_model,
        )
        logger.info("Indexed %d chunks from %s", len(self._chunks), self._filename)
        return len(self._chunks)

    # ── Retrieval ───────────────────────────────────────────────────────────
    def search(self, query: str, k: int = 5) -> List[Document]:
        if not self.is_ready:
            return []
        normalized_query = normalize_upload_query(query)
        semantic = self._vectorstore.similarity_search(normalized_query, k=k)

        # Dense retrieval helps paraphrases, but exact document headings can
        # be missed when a user writes them without spaces. Merge in lexical
        # matches from the entire uploaded document, then de-duplicate.
        query_terms = _terms(query)
        lexical = sorted(
            self._chunks,
            key=lambda chunk: len(query_terms & _terms(chunk.page_content)),
            reverse=True,
        )
        lexical_matches = [
            chunk for chunk in lexical
            if query_terms & _terms(chunk.page_content)
        ]
        results: List[Document] = []
        seen: set[str] = set()
        for document in [*lexical_matches, *semantic]:
            fingerprint = re.sub(r"\s+", " ", document.page_content).strip().casefold()
            if fingerprint and fingerprint not in seen:
                seen.add(fingerprint)
                results.append(document)
            if len(results) >= k:
                break
        return results

    def search_with_scores(self, query: str, k: int = 5) -> List[tuple]:
        if not self.is_ready:
            return []
        return self._vectorstore.similarity_search_with_score(query, k=k)

    def clear(self) -> None:
        self._vectorstore = None
        self._chunks = []
        self._filename = ""
